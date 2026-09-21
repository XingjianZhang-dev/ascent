#!/usr/bin/env python3
"""Create the separate Elsevier declaration-of-interest DOCX."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_font(run, *, size: float = 11, bold: bool = False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold


def add_labeled_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    set_font(paragraph.add_run(f"{label}: "), bold=True)
    set_font(paragraph.add_run(value))


def build(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    set_font(title.add_run("Declaration of Interests"), size=16, bold=True)

    add_labeled_paragraph(
        document,
        "Manuscript title",
        "ASCENT: Scale-Complementary External State for Frozen Long-Context Language Models",
    )
    add_labeled_paragraph(document, "Journal", "Array")
    add_labeled_paragraph(document, "Author", "Xingjian Zhang")

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(8)
    set_font(heading.add_run("Declaration"), size=13, bold=True)

    statement = document.add_paragraph()
    statement.paragraph_format.space_before = Pt(0)
    statement.paragraph_format.space_after = Pt(18)
    statement.paragraph_format.line_spacing = 1.10
    set_font(
        statement.add_run(
            "The author declares that he has no known competing financial "
            "interests or personal relationships that could have appeared to "
            "influence the work reported in this paper."
        )
    )

    add_labeled_paragraph(document, "Corresponding author", "Xingjian Zhang")
    add_labeled_paragraph(document, "Email", "xingjiaz@andrew.cmu.edu")
    add_labeled_paragraph(document, "Date", "August 23, 2026")

    document.core_properties.title = "Declaration of Interests"
    document.core_properties.subject = "Array submission"
    document.core_properties.author = "Xingjian Zhang"
    document.core_properties.keywords = "declaration of interests"
    document.core_properties.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
