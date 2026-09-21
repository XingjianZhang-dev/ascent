#!/usr/bin/env python3
"""Build the separate author-identifying title page for Array."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "ASCENT: Scale-Complementary External State for Frozen Long-Context Language Models"
INK = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)


def set_font(run, size: float, *, bold: bool = False, italic: bool = False, color=None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph(paragraph, *, before=0, after=6, line=1.10, align=None) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_labeled(doc: Document, label: str, value: str, *, after=3) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, after=after)
    set_font(p.add_run(f"{label}: "), 10.5, bold=True)
    set_font(p.add_run(value), 10.5)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, before=8, after=3)
    set_font(p.add_run(text), 12.5, bold=True, color=INK)


def build(output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    set_paragraph(header, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(header.add_run("ARRAY | REGULAR PAPER: RESEARCH PAPER"), 9, bold=True, color=MUTED)

    p = doc.add_paragraph()
    set_paragraph(p, before=10, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("TITLE PAGE"), 12, bold=True, color=INK)

    p = doc.add_paragraph()
    set_paragraph(p, after=10, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run(TITLE), 19, bold=True)

    p = doc.add_paragraph()
    set_paragraph(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run("Xingjian Zhang"), 14, bold=True)

    p = doc.add_paragraph()
    set_paragraph(p, after=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(
        p.add_run("Department of Electrical and Computer Engineering, Carnegie Mellon University"),
        10.5,
        italic=True,
        color=MUTED,
    )

    add_heading(doc, "Corresponding author")
    add_labeled(doc, "Name", "Xingjian Zhang")
    add_labeled(
        doc,
        "Address",
        "Department of Electrical and Computer Engineering, Carnegie Mellon University, "
        "5000 Forbes Avenue, Pittsburgh, Pennsylvania 15213, USA",
    )
    add_labeled(doc, "Email", "xingjiaz@andrew.cmu.edu")
    add_labeled(doc, "ORCID", "https://orcid.org/0009-0009-9141-8819")
    add_labeled(doc, "Telephone", "+1 917 940 6399; +86 139 5192 4078")

    add_heading(doc, "Acknowledgements")
    p = doc.add_paragraph()
    set_paragraph(p, after=3)
    set_font(p.add_run("None."), 10.5)

    add_heading(doc, "Funding")
    p = doc.add_paragraph()
    set_paragraph(p, after=3)
    set_font(
        p.add_run(
            "This research did not receive any specific grant from funding agencies in the "
            "public, commercial, or not-for-profit sectors."
        ),
        10.5,
    )

    add_heading(doc, "Declaration of interests")
    p = doc.add_paragraph()
    set_paragraph(p, after=3)
    set_font(
        p.add_run(
            "The author declares that he has no known competing financial interests or personal "
            "relationships that could have appeared to influence the work reported in this paper."
        ),
        10.5,
    )

    add_heading(doc, "CRediT authorship contribution statement")
    p = doc.add_paragraph()
    set_paragraph(p, after=0)
    set_font(p.add_run("Xingjian Zhang: "), 10.5, bold=True)
    set_font(
        p.add_run(
            "Conceptualization, Methodology, Software, Formal analysis, Investigation, "
            "Validation, Data curation, Visualization, Writing - original draft, Writing - "
            "review and editing, Project administration."
        ),
        10.5,
    )

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Author-identifying title page for Array"
    doc.core_properties.author = "Xingjian Zhang"
    doc.core_properties.keywords = "Array; title page; ASCENT"
    doc.core_properties.comments = "Separate from the double-anonymized manuscript"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("paper/submission/Title_Page.docx"),
    )
    args = parser.parse_args()
    build(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
